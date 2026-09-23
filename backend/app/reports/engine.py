"""Render the institutional report from an analysis.

Threat model (docs/threat-model.md → Finding snippets): every value that came
out of the audited tree — paths, snippets, messages, package names — is hostile.

- HTML: Jinja2 with autoescape forced on for every template, ``StrictUndefined``
  so a typo in a template fails loudly instead of printing nothing.
- Markdown: a dedicated ``md_escape`` filter; snippets are fenced with a fence
  longer than any backtick run inside them, so content can never close the
  fence early.
- PDF: WeasyPrint receives the escaped HTML and a ``url_fetcher`` that ONLY
  serves ``file://`` URLs under this package's own template directory. A hostile
  snippet that survived escaping as a URL would still be refused: the renderer
  is never an HTTP client. No external fonts (Hard Rule: no CDN).
- DOCX: python-docx writes text runs; there is no markup channel to escape.
"""

from __future__ import annotations

import io
import logging
import re
from collections.abc import Sequence
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

from jinja2 import Environment, FileSystemLoader, StrictUndefined, select_autoescape
from markupsafe import Markup, escape

from app.analysis.models import Analysis
from app.projects.models import Project
from app.reports.context import build_context
from app.reports.errors import ForbiddenAssetFetch, ReportRenderError
from app.reports.models import ReportVersion
from app.reports.strings import report_strings
from app.workflow.triage import strip_control_chars

logger = logging.getLogger("dioptra.reports")

TEMPLATES_DIR = (Path(__file__).resolve().parents[2] / "templates" / "report").resolve()
HTML_TEMPLATE = "report.html.j2"
MARKDOWN_TEMPLATE = "report.md.j2"

_BACKTICK_RUN = re.compile(r"`+")
# Characters that can open Markdown structure inline: backslash, backtick,
# emphasis, links, table cells. Intraword underscores (snake_case paths and
# package names) are left alone — CommonMark does not treat them as emphasis,
# and escaping them makes every path unreadable.
_MD_SPECIAL = re.compile(r"([\\`*\[\]|])")
_MD_UNDERSCORE = re.compile(r"(?<![A-Za-z0-9])_|_(?![A-Za-z0-9])")
_MD_LINE_START = re.compile(r"^([#>+\-])", re.MULTILINE)


@dataclass(frozen=True)
class ReportBundle:
    html: str
    markdown: str


def md_escape(value: object) -> Markup:
    """Neutralize Markdown and HTML syntax in an inline value.

    Returns ``Markup`` because the result is final: the environment's autoescape
    must not re-escape the entities this filter already produced.
    """
    text = "" if value is None else str(value)
    text = _MD_SPECIAL.sub(r"\\\1", text)
    text = _MD_UNDERSCORE.sub(r"\\_", text)
    text = _MD_LINE_START.sub(r"\\\1", text)
    # Backslash-escaping `<` is not honoured by every Markdown renderer; the
    # entity form is, so a raw tag can never re-form.
    return escape(text)


def md_fence(value: object) -> Markup:
    """Wrap a code snippet in a fence longer than any backtick run it contains."""
    text = "" if value is None else str(value)
    longest = max((len(run) for run in _BACKTICK_RUN.findall(text)), default=0)
    fence = Markup("`") * max(3, longest + 1)
    newline = Markup("\n")
    return fence + newline + escape(text) + newline + fence


def html_pre(value: object) -> Markup:
    """Escape a snippet for a ``<pre>`` block; ``Markup.format`` escapes the value."""
    return Markup('<pre class="snippet">{}</pre>').format("" if value is None else str(value))


@lru_cache(maxsize=1)
def stylesheet() -> str:
    """The report stylesheet, read once.

    It is INLINED into the HTML: the API serves this HTML to browsers under a
    ``style-src 'self'`` CSP, and the PDF path needs no fetch at all. The
    restricted fetcher stays as the guard for anything WeasyPrint might still
    try to resolve.
    """
    return (TEMPLATES_DIR / "report.css").read_text(encoding="utf-8")


@lru_cache(maxsize=1)
def _html_environment() -> Environment:
    env = Environment(
        loader=FileSystemLoader(str(TEMPLATES_DIR)),
        autoescape=select_autoescape(default=True, default_for_string=True),
        undefined=StrictUndefined,
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.filters["pre"] = html_pre
    return env


@lru_cache(maxsize=1)
def _markdown_environment() -> Environment:
    # Autoescape stays ON even here: any value not passed through ``md_escape``
    # or ``md_fence`` is still HTML-escaped, never raw.
    env = Environment(
        loader=FileSystemLoader(str(TEMPLATES_DIR)),
        autoescape=select_autoescape(default=True, default_for_string=True),
        undefined=StrictUndefined,
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.filters["md"] = md_escape
    env.filters["fence"] = md_fence
    return env


def _allowed_asset(url: str) -> Path:
    """The on-disk file a URL may resolve to, or :class:`ForbiddenAssetFetch`."""
    parsed = urlparse(url)
    if parsed.scheme != "file" or parsed.netloc not in ("", "localhost"):
        raise ForbiddenAssetFetch(f"refused non-file asset: {parsed.scheme or '(none)'}")
    candidate = Path(unquote(parsed.path)).resolve()
    if not candidate.is_relative_to(TEMPLATES_DIR) or not candidate.is_file():
        raise ForbiddenAssetFetch("refused asset outside the report templates")
    return candidate


def _restricted_fetcher_class() -> type[Any]:
    """Build the WeasyPrint fetcher lazily so the module imports without pango."""
    from weasyprint.urls import URLFetcher  # type: ignore[import-untyped]  # noqa: PLC0415

    class RestrictedURLFetcher(URLFetcher):  # type: ignore[misc]
        """Serves only files under the report template directory.

        ``fail_on_errors=True``: a refused reference aborts the PDF instead of
        producing a document silently missing its own stylesheet.
        """

        def __init__(self) -> None:
            super().__init__(allowed_protocols={"file"}, allow_redirects=False, fail_on_errors=True)

        def fetch(self, url: str, headers: dict[str, str] | None = None) -> Any:
            return super().fetch(_allowed_asset(url).as_uri(), headers)

    return RestrictedURLFetcher


def restricted_url_fetcher(url: str) -> Any:
    """Fetch one template asset through the restricted fetcher (test/diagnostic entry)."""
    return _restricted_fetcher_class()()(url)


def render_html(
    analysis: Analysis,
    project: Project,
    *,
    version: ReportVersion | None = None,
    versions: Sequence[ReportVersion] = (),
) -> str:
    context = build_context(analysis, project, version=version, versions=versions)
    template = _html_environment().get_template(HTML_TEMPLATE)
    # The stylesheet is a trusted, static file of ours: it is the one value that
    # legitimately bypasses autoescape.
    return template.render(**context, stylesheet=Markup(stylesheet()))  # noqa: S704


def render_markdown(
    analysis: Analysis,
    project: Project,
    *,
    version: ReportVersion | None = None,
    versions: Sequence[ReportVersion] = (),
) -> str:
    context = build_context(analysis, project, version=version, versions=versions)
    template = _markdown_environment().get_template(MARKDOWN_TEMPLATE)
    return template.render(**context)


def render_bundle(analysis: Analysis, project: Project) -> ReportBundle:
    return ReportBundle(
        html=render_html(analysis, project), markdown=render_markdown(analysis, project)
    )


def render_pdf(
    analysis: Analysis,
    project: Project,
    *,
    version: ReportVersion | None = None,
    versions: Sequence[ReportVersion] = (),
) -> bytes:
    html = render_html(analysis, project, version=version, versions=versions)
    try:
        from weasyprint import HTML  # type: ignore[import-untyped]  # noqa: PLC0415

        fetcher = _restricted_fetcher_class()()
        document = HTML(string=html, base_url=TEMPLATES_DIR.as_uri() + "/", url_fetcher=fetcher)
        pdf = document.write_pdf()
    except BaseException as exc:
        # WeasyPrint's FatalURLFetchingError derives from BaseException, so a
        # refused fetch would otherwise escape every handler as a bodiless 500.
        if isinstance(exc, KeyboardInterrupt | SystemExit):
            raise
        logger.exception("pdf rendering failed for analysis %s", analysis.id)
        raise ReportRenderError("weasyprint failed") from exc
    if not isinstance(pdf, bytes):
        raise ReportRenderError("weasyprint returned no bytes")
    return pdf


def _docx_table(document: Any, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = strip_control_chars(str(value))


def _na(value: object, suffix: str = "") -> str:
    return "N/A" if value is None else f"{value}{suffix}"


def _docx_closure(doc: Any, context: dict[str, Any], strings: dict[str, Any]) -> None:
    """Sections 7–10 as headings, paragraphs and tables (text runs only)."""

    def add(text: object = "", **kwargs: Any) -> Any:
        return doc.add_paragraph(strip_control_chars(str(text)), **kwargs)

    heading = strings["headings"]
    metrics, debt, inventory, annexes = (
        context["metrics"],
        context["test_debt"],
        context["inventory"],
        context["annexes"],
    )
    sm, sd, si, sa = (
        strings["metrics"],
        strings["test_debt"],
        strings["inventory"],
        strings["annexes"],
    )

    doc.add_heading(heading["metrics"], level=1)
    add(sm["intro"])
    if metrics["functions_measured"]:
        add(
            sm["summary"].format(
                functions=metrics["functions_measured"],
                complex=metrics["complex_functions"],
                commented=metrics["commented_files"],
            )
        )
        _docx_table(
            doc,
            [
                sm["labels"]["ccn"],
                sm["labels"]["nloc"],
                sm["labels"]["function"],
                sm["labels"]["path"],
            ],
            [[str(f["ccn"]), str(f["nloc"]), f["function"], f["path"]] for f in metrics["top"]],
        )
        add(sm["duplication"])
    else:
        add(sm["none"])

    doc.add_heading(heading["test_debt"], level=1)
    add(sd["intro"])
    add(sd["stage_note"].format(stage=context["stage_label"]))
    if debt["plan"]:
        totals = debt["totals"]
        add(
            sd["plan_line"].format(
                functions=totals["functions"],
                criterion=debt["plan"]["criterion_label"],
                author=debt["plan"]["created_by"],
            )
        )
        add(debt["plan"]["rationale"])
        add(
            sd["totals"].format(
                approved=totals["approved"],
                functions=totals["functions"],
                written=totals["written"],
                passed=totals["passed"],
                statements=_na(totals["statement_percent"], " %"),
                branches=_na(totals["branch_percent"], " %"),
                survivors=totals["survivors"],
                equivalent=totals["equivalent"],
            )
        )
        if totals.get("mutation_not_measured"):
            # Same sentence the HTML and Markdown carry: "mutantes vivos: 0"
            # under the totals would otherwise read as a measured zero.
            add(sd["not_measured_note"].format(count=totals["mutation_not_measured"]))
        lbl = sd["labels"]
        _docx_table(
            doc,
            [
                lbl["function"],
                lbl["cases"],
                lbl["written"],
                lbl["statements"],
                lbl["branches"],
                lbl["status"],
                lbl["survivors"],
            ],
            [
                [
                    f"{f['path']} · {f['function']}",
                    str(f["cases"]),
                    lbl["yes"] if f["written"] else lbl["no"],
                    _na(f["statement_percent"], " %"),
                    _na(f["branch_percent"], " %"),
                    f["status_label"],
                    # "no medida" instead of 0 when the tool could not mutate
                    # this function at all (a free PHP function under
                    # Infection): zero survivors and no measurement are not
                    # the same fact. tasks/phase7a-php.md.
                    str(f["survivors"])
                    if f.get("mutation_measured", True)
                    else lbl["not_measured"],
                ]
                for f in debt["functions"]
            ],
        )
    else:
        add(sd["no_plan"])

    doc.add_heading(heading["inventory"], level=1)
    add(si["intro"])
    if inventory is None:
        add(si["no_sbom"])
    else:
        add(
            si["summary"].format(
                components=inventory["components"],
                licenses=", ".join(row["name"] for row in inventory["licenses"]) or "N/A",
                unlicensed=inventory["unlicensed"],
            )
        )
        if not inventory["correlated"]:
            add(si["not_correlated"])
        else:
            add(
                si["vulndb_line"].format(date=inventory["vulndb_last_update"])
                if inventory["vulndb_last_update"]
                else si["vulndb_none"]
            )
            doc.add_heading(si["open_title"], level=2)
            lbl = si["labels"]
            if inventory["open"]:
                _docx_table(
                    doc,
                    [
                        lbl["component"],
                        lbl["version"],
                        lbl["cve"],
                        lbl["severity"],
                        lbl["fixed_in"],
                        lbl["vex"],
                    ],
                    [
                        [
                            r["component"],
                            r["version"],
                            r["cve"],
                            r["severity_label"],
                            r["fixed_in"],
                            r["vex_label"],
                        ]
                        for r in inventory["open"]
                    ],
                )
            else:
                add(si["open_none"])
            if inventory["not_affected"]:
                doc.add_heading(si["not_affected_title"], level=2)
                for r in inventory["not_affected"]:
                    add(
                        f"{r['component']} {r['version']} — {r['cve']}: {r['justification']}",
                        style="List Bullet",
                    )
            if inventory["outdated"]:
                doc.add_heading(si["outdated_title"], level=2)
                add(si["outdated_note"])
                _docx_table(
                    doc,
                    [lbl["component"], lbl["version"], lbl["newer"]],
                    [[r["component"], r["version"], r["newer"]] for r in inventory["outdated"]],
                )
            doc.add_heading(si["crypto_title"], level=2)
            if inventory["crypto"]:
                _docx_table(
                    doc,
                    [lbl["algorithm"], lbl["primitive"], lbl["location"]],
                    [
                        [
                            r["algorithm"] + (f" — {si['weak']}" if r["weak"] else ""),
                            r["primitive"],
                            r["location"],
                        ]
                        for r in inventory["crypto"]
                    ],
                )
            else:
                add(si["crypto_none"])

    doc.add_heading(heading["annexes"], level=1)
    doc.add_heading(heading["annex_asvs"], level=2)
    add(sa["asvs_intro"])
    _docx_table(
        doc,
        [sa["asvs_labels"]["chapter"], sa["asvs_labels"]["title"], sa["asvs_labels"]["findings"]],
        [[r["chapter"], r["title"], str(r["findings"])] for r in annexes["asvs"]],
    )
    doc.add_heading(heading["annex_diagrams"], level=2)
    add(sa["diagrams_intro"])
    for d in annexes["diagrams"]:
        doc.add_heading(f"{d['path']} · {d['function']}", level=3)
        # DOCX carries no SVG channel here: the Mermaid interchange text is the diagram.
        add(d["mermaid"] if d["mermaid"] else sa["diagram_unavailable"], style="No Spacing")
    if not annexes["diagrams"]:
        add(sa["cases_none"])
    doc.add_heading(heading["annex_cases"], level=2)
    add(sa["cases_intro"])
    for group in annexes["cases"]:
        state = sa["approved"] if group["approved"] else sa["not_approved"]
        doc.add_heading(f"{group['path']} · {group['function']} ({state})", level=3)
        for case in group["cases"]:
            covers = f" [{', '.join(case['covers'])}]" if case["covers"] else ""
            add(f"{case['id']} — {case['title']}{covers}", style="List Bullet")
    if not annexes["cases"]:
        add(sa["cases_none"])
    doc.add_heading(heading["annex_tests"], level=2)
    add(sa["tests_intro"])
    for t in annexes["tests"]:
        doc.add_heading(t["filename"], level=3)
        add(t["content"], style="No Spacing")
        if t["truncated"]:
            add(sa["truncated"])
    if not annexes["tests"]:
        add(sa["tests_none"])
    doc.add_heading(heading["annex_bom"], level=2)
    add(sa["bom_intro"])
    if annexes["sbom"]:
        add(
            sa["bom_line"].format(
                components=annexes["sbom"]["components"],
                generator=annexes["sbom"]["generator"],
                spec=annexes["sbom"]["spec_version"],
            )
        )
    else:
        add(si["no_sbom"])


def render_docx(
    analysis: Analysis,
    project: Project,
    *,
    version: ReportVersion | None = None,
    versions: Sequence[ReportVersion] = (),
) -> bytes:
    """Basic DOCX: headings, paragraphs and the tables. Text runs only.

    Every label comes from ``templates/report/strings.json``; the prose is the
    rendered version's sections (defaults or the analyst's overrides).
    """
    context = build_context(analysis, project, version=version, versions=versions)
    try:
        from docx import Document  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ReportRenderError("python-docx unavailable") from exc

    strings = report_strings()
    heading = strings["headings"]
    label = strings["labels"]
    doc = Document()

    def add(text: object = "", **kwargs: Any) -> Any:
        # python-docx raises on XML-incompatible control characters, and the
        # audited tree can put one anywhere (a path, a component name).
        return doc.add_paragraph(strip_control_chars(str(text)), **kwargs)

    doc.add_heading(strings["cover_title"], level=0)
    add(context["system"]["name"])
    add(context["period"])
    doc.add_page_break()  # type: ignore[no-untyped-call]

    sections = context["sections"]
    distribution = strings["distribution"]
    doc.add_heading(heading["introduction"], level=1)
    for paragraph in sections["introduction"]:
        add(paragraph)
    doc.add_heading(heading["summary"], level=1)
    for paragraph in sections["summary"]:
        add(paragraph)
    if context["summary"]["total"]:
        _docx_table(
            doc,
            [distribution["severity"], distribution["count"]],
            [[r["label"], str(r["count"])] for r in context["summary"]["by_severity"]],
        )
        _docx_table(
            doc,
            [distribution["category"], distribution["count"]],
            [[r["label"], str(r["count"])] for r in context["summary"]["by_owasp"]],
        )

    doc.add_heading(heading["system"], level=1)
    system = context["system"]
    _docx_table(
        doc,
        [label["field"], label["value"]],
        [
            [label["site_name"], system["name"]],
            [label["framework"], system["framework"]],
            [label["installed_at"], system["installed_at"]],
            [label["database"], system["database"]],
            [label["developer"], system["developer"]],
        ],
    )
    doc.add_heading(heading["version_control"], level=2)
    _docx_table(
        doc,
        [label["version"], label["areas"], label["change"], label["delivered_at"]],
        [
            [r["version"], r["areas"], r["description"], r["delivered_at"]]
            for r in context["version_control"]
        ],
    )

    def finding_block(finding: dict[str, Any]) -> None:
        doc.add_heading(finding["title"], level=2)
        _docx_table(
            doc,
            [finding["title"], label["severity_level"], finding["severity_label"]],
            [[finding["subtitle"], label["owasp_code"], finding["owasp_label"]]],
        )
        doc.add_heading(heading["description"], level=3)
        add(finding["description"])
        if finding["message"]:
            add(finding["message"])
        doc.add_heading(heading["impact"], level=3)
        add(finding["impact"])
        doc.add_heading(heading["detection"], level=3)
        add(heading["path"])
        add(finding["location"])
        if finding["snippet"]:
            add(finding["snippet"], style="No Spacing")
        doc.add_heading(heading["mitigation"], level=3)
        for item in finding["mitigation"]:
            add(item, style="List Bullet")
        if finding["references"]:
            doc.add_heading(heading["references"], level=3)
            for item in finding["references"]:
                add(item, style="List Bullet")

    if context["security_findings"]:
        doc.add_heading(heading["findings"], level=1)
        for paragraph in sections["findings_intro"]:
            add(paragraph)
        for finding in context["security_findings"]:
            finding_block(finding)

    doc.add_heading(heading["dependencies"], level=1)
    if context["dependency_findings"] or context["dependency_scan_ran"]:
        for paragraph in sections["dependencies_intro"]:
            add(paragraph)
        if context["dependency_findings"]:
            add(label["found_below"])
            for finding in context["dependency_findings"]:
                finding_block(finding)
        else:
            add(strings["dependencies_clean"])
    else:
        add(strings["dependencies_not_run"])

    if context["commented_code_files"]:
        doc.add_heading(heading["practices"], level=1)
        doc.add_heading(heading["commented_code"], level=2)
        for paragraph in sections["practices"]:
            add(paragraph)
        add(label["commented_files"])
        for path in context["commented_code_files"]:
            add(path, style="List Bullet")

    doc.add_heading(heading["coverage"], level=1)
    for paragraph in sections["coverage_intro"]:
        add(paragraph)
    _docx_table(
        doc,
        [label["tool"], label["category"], label["status"], label["detail"]],
        [[r["tool"], r["category"], r["status_label"], r["detail"]] for r in context["tool_runs"]],
    )
    _docx_closure(doc, context, strings)
    add(f"{label['author_prefix']} {context['author']}")

    buffer = io.BytesIO()
    try:
        doc.save(buffer)
    except Exception as exc:  # pragma: no cover - python-docx is deterministic
        raise ReportRenderError("docx save failed") from exc
    return buffer.getvalue()


def report_file_name(project_name: str, fmt: str) -> str:
    """Download name for an export; the stem is a plain header-safe token."""
    strings = report_strings()
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", project_name).strip("-")[:60]
    return f"{strings['filename']['prefix']}{stem or strings['filename']['fallback']}.{fmt}"
